#!/usr/bin/env python3
"""Generate C-level Claude + Cursor training deck (PPTX). Run from repo root."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

OUT = Path(__file__).resolve().parents[1] / "06-Resources" / "C_Level_Claude_Cursor_Training_Deck.pptx"
PROFILE_PHOTO = Path(__file__).resolve().parents[1] / "06-Resources" / "profile-photo.png"

SLIDES = [
    (
        "AI CMS - Программа развития AI-лидерства",
        [
            "Подготовил: Roman Matsukatov",
            "Дата: 08 апреля 2026",
        ],
    ),
    (
        "Обо мне",
        [
            "12+ лет в Product Management на позициях Senior/Lead/CPO.",
            "Запуск AI-продуктов с нуля.",
            "Внедрение Agile-практик в командах.",
            "Опыт запуска Compliance-направления в iGaming (UKGC, MGA, Curacao, Ontario, Anjouan, Tobique, New Jersey).",
            "Для вашей компании это обучение с фокусом на внедрение, governance и измеримый бизнес-результат.",
        ],
    ),
    (
        "Целевая модель AI CMS: недели 1-2",
        [
            "• C-level alignment и единый словарь решений.",
            "• Классы данных и правила использования AI по уровням риска.",
            "• Допустимые AI-среды и сценарии для руководителей.",
            "• Адаптивный AI workflow для C-level: от приоритизации реальных кейсов к единому циклу анализа, риск-контроля, решения и внедрения.",
        ],
    ),
    (
        "Целевая модель AI CMS: недели 3-4",
        [
            "• Запуск HR и Support воркфлоу на реальных кейсах.",
            "• Короткие демо 2-5 минут в рабочем ритме команд.",
            "• Связка с календарем, почтой, чатами и базой знаний.",
            "• Фиксация обратной связи и корректировка playbook по итогам недели.",
        ],
    ),
    (
        "Целевая модель AI CMS: месяц 1-2",
        [
            "• Масштабирование в Finance/Ops после подтверждения первых кейсов.",
            "• Owner'ы по воркфлоу в каждом направлении бизнеса.",
            "• KPI adoption и качества на регулярном review.",
            "• Эскалация спорных кейсов по данным/compliance через Program Owner и Risk Owner.",
        ],
    ),
    (
        "Волна 1: C-level alignment",
        [
            "Фиксируем классификацию данных, бюджетные лимиты, модель рисков и эскалаций.",
            "На выходе, единая рамка, утверждённые владельцы и критерии успеха первой волны.",
        ],
    ),
    (
        "Волна 2: функциональные треки",
        [
            "Каждая функция получает сценарии на своих процессах, с общим стандартом качества и безопасности.",
            "На выходе, руководители функций самостоятельно проходят типовой кейс и подтверждают готовность команды.",
        ],
    ),
    (
        "Волна 3: пилоты и масштабирование",
        [
            "Запускаем 1-2 пилота со сквозной ответственностью и измеримым бизнес-эффектом.",
            "Тиражируем только подтверждённые практики, что даёт масштаб без хаоса.",
        ],
    ),
    (
        "KPI и контроль эффекта",
        [
            "Метрики: скорость решений, время на типовые задачи, качество артефактов, adoption, риск-инциденты.",
            "C-level смотрит на результат по KPI, а не на количество экспериментов или инструментов.",
        ],
    ),
    (
        "Операционная модель: кто за что отвечает",
        [
            "Sponsor (C-level), Program Owner, Function Leads, Risk Owner, с четкой зоной ответственности каждого.",
            "Cadence: еженедельный операционный ритм и ежемесячный управленческий review с решениями по масштабу.",
        ],
    ),
    (
        "Риски и контрмеры",
        [
            "Риски: утечки данных, неуправляемые расходы, низкий adoption, локальная автоматизация без эффекта для бизнеса.",
            "Контрмеры: policy-by-default, бюджетные лимиты, KPI-гейты, обязательная эскалация спорных кейсов.",
        ],
    ),
    (
        "Первые 30 дней: план запуска",
        [
            "Неделя 1: утвердить рамку и владельцев. Неделя 2: запустить C-level alignment сессию.",
            "Недели 3-4: старт функциональных треков, сбор baseline KPI, запуск первого пилота.",
        ],
    ),
    (
        "Решение, которое нужно принять сегодня",
        [
            "Утвердить программу на 3 волны, назначить владельцев и зафиксировать календарь первой волны.",
            "Подтвердить KPI и формат ежемесячного управленческого review по внедрению AI.",
        ],
    ),
    (
        "AI Literacy и AI-Thinking — дополнительный модуль",
        [
            "Зачем: ясная цель и собранный контекст, согласованный план и предсказуемый цикл до результата, а не бесконечные разовые чаты.",
            "Фокус: откуда берётся контекст (тулы, чаты, внешние системы), планирование до имплементации, цикл со сверкой и апрувом, повторяемое оформляем в артефакты, чтобы не платить токены за одни и те же факты в каждом запросе.",
        ],
    ),
    (
        "Что разбираем в модуле",
        [
            "Цель: зафиксировать, что делаем и зачем.",
            "Контекст: собрать релевантную информацию в одном месте; для этого подключаем нужные тулы, чаты и внешние системы.",
            "Планирование: сначала описание намерения и шагов, без немедленного прыжка в имплементацию.",
            "Цикл с AI: план → сверка → апрув → выполнение → уточнение и правки → финальная версия.",
            "Повторяемое превращаем в артефакты: шаблоны, инструкции, база знаний — сильнее контекст, меньше токенов на повтор одних и тех же фактов.",
        ],
    ),
]


def add_title_slide(prs, title, subtitle_lines, include_photo_block=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    left, top, width, height = Inches(0.6), Inches(1.0), Inches(9.0), Inches(1.2)
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    body_top = Inches(2.4)
    body_width = Inches(5.8) if include_photo_block else Inches(9.0)
    body = slide.shapes.add_textbox(left, body_top, body_width, Inches(4.5))
    btf = body.text_frame
    btf.clear()
    for i, line in enumerate(subtitle_lines):
        para = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
        para.text = line
        para.font.size = Pt(18)
        para.space_after = Pt(10)
        para.level = 0
        para.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    if include_photo_block:
        photo_left = Inches(6.7)
        photo_top = Inches(2.1)
        photo_w = Inches(2.6)
        photo_h = Inches(3.2)
        if PROFILE_PHOTO.exists():
            slide.shapes.add_picture(str(PROFILE_PHOTO), photo_left, photo_top, photo_w, photo_h)
        else:
            photo = slide.shapes.add_shape(
                1,  # MSO_SHAPE.RECTANGLE
                photo_left,
                photo_top,
                photo_w,
                photo_h,
            )
            photo.fill.solid()
            photo.fill.fore_color.rgb = RGBColor(0xF1, 0xF3, 0xF5)
            photo.line.color.rgb = RGBColor(0xC7, 0xCD, 0xD6)
            ptf = photo.text_frame
            ptf.clear()
            pp = ptf.paragraphs[0]
            pp.text = "Your photo"
            pp.font.size = Pt(18)
            pp.font.bold = True
            pp.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    for idx, (title, bullets) in enumerate(SLIDES):
        add_title_slide(prs, title, bullets, include_photo_block=(idx == 0))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
