#!/usr/bin/env python3
"""Write flat .txt previews under .scripts/.cache/ (not in the bundle — no 03-extracted-previews in repo or on Drive)."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
import openpyxl

REPO = Path(__file__).resolve().parent.parent
BUNDLE = (
    REPO
    / "04-Projects"
    / "Banda"
    / "banda-drive-import"
    / "2. Banda Presentations Bundle"
)
# Вне vault-пакета: только локальный кэш для поиска по тексту (gitignore).
OUT = REPO / ".scripts" / ".cache" / "banda-presentations-previews"


def docx_text(p: Path) -> str:
    d = Document(p)
    return "\n".join(x.text for x in d.paragraphs if x.text.strip())


def xlsx_preview(p: Path, max_rows: int = 100) -> str:
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    parts = []
    for sn in wb.sheetnames[:8]:
        ws = wb[sn]
        rows = []
        for _, row in enumerate(ws.iter_rows(max_row=max_rows, values_only=True)):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                rows.append(" | ".join(cells[:18]))
        parts.append(f"--- sheet: {sn} ---\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts)


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    for sub, pat in [
        ("New Primary Slide decks (Drafts)", "*.docx"),
        ("Secondary Slide decks (Drafts)", "*.xlsx"),
    ]:
        d = BUNDLE / sub
        if not d.is_dir():
            print("skip missing", d, file=sys.stderr)
            continue
        for f in sorted(d.glob(pat)):
            stem = f.stem.replace("/", "-")
            try:
                if f.suffix == ".docx":
                    t = docx_text(f)
                else:
                    t = xlsx_preview(f)
                (OUT / f"{sub}__{stem}.txt").write_text(t, encoding="utf-8")
                print(stem, len(t))
            except Exception as e:
                print("ERR", f, e, file=sys.stderr)
                return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
